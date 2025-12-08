# 入参文件地址，工作簿名称，表格1，表格2，输出表格（可选）；
# 讲表格2的str内容拼接到表格1的str内容上去
import openpyxl as xl
def cell_value_concatenator(file_path: str,
                            work_sheet: str,
                            cell1: str,
                            cell2: str,
                            output_cell: str = ''):
    wb = xl.load_workbook(file_path)
    ws = wb[work_sheet]
    val1 = str(ws[cell1].value)
    val2 = str(ws[cell2].value)
    if output_cell:
        ws[output_cell].value = val1 + val2
    return val1 + val2

if __name__ == '__main__':
    file_path = "/Users/edison/Desktop/J716 S2F/DVT2/FATP/C/stuff 练手/PWZM DVT2c Build Rules V5.xlsx"
    work_sheet = 'HVE 11_7 - J716C DVT-2 Build ru'
    cell1 = 'D18'
    for row in range(3,15):
        cell2 = f'B{row}'
        print(cell2)
        print(cell_value_concatenator(file_path, work_sheet, cell1, cell2))
